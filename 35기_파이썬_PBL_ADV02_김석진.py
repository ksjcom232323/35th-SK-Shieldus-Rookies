import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from datetime import datetime
import os
import smtplib
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from dotenv import load_dotenv
from pathlib import Path
from html import escape
import traceback

load_dotenv()

def get_security_news():
    '''
    "https://www.boannews.com/"에서 뉴스 수집함수

    Return : news_list(수집된 뉴스리스트)
    '''
    url="https://www.boannews.com/"

    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        )
    }


    try :
        # 웹페이지 요청
        response = requests.get(
            url,
            headers= headers,
            timeout=10
        )

        response.raise_for_status()
        print("연결 성공")
        print("상태 코드:",response.status_code)

        #html 파싱
        soup = BeautifulSoup(response.text, "html.parser")

        links = soup.find_all("a")
        news_list =[]
        seen_titles = set()
        risk_titles = ('취약점','유출','해킹')
        for a in links:

            # 제목
            title = a.get_text(strip=True) # 문자열 공백제거
            # HTML에서 텍스트를 가져올 때 불필요한 앞뒤 공백이나 줄바꿈을 제거해서 깔끔하게 가져오라는 의미

            # 링크 주소
            href = a.get("href")

            # 제목이나 링크가 없으면 제외
            if not title or not href:  continue

            # 너무 짧은 메뉴 제외
            if len(title) < 15: continue

            # 중복 제목 제외
            if title in seen_titles:  continue

            # 전체 URL 만들기
            link = urljoin(url, href)

            # 보안뉴스 외부 링크 제외
            if url not in link:  continue

            seen_titles.add(title)

            #risk 등급 부여
            risk = "High" if any(word in title for word in risk_titles) else "Normal"

            #news_list 생성
            news_list.append({"title":title,"link":link,"risk":risk})

        #수집된 뉴스가 5건미만이면 종료
        if len(news_list) < 5:
            print(f"❌ 수집된 뉴스 기사가 5건미만이면 프로그램을 종료합니다. 수집된 뉴스는 {len(news_list)}건입니다. ")
            raise SystemExit(1)
        
    

        print("뉴스 가져오기 성공")
        return news_list 

    except requests.exceptions.RequestException as e:
        print("❌ 웹사이트 연결 오류")
        print(e)

def create_excel_report (news_list,date):
    '''
    수집된 뉴스를 엑셀파일로 저장
    
    Args : news_list(수집된 뉴스리스트)
            date(날짜)
    Return : excel_path(생성된 엑셀주소)
    '''
    wb = Workbook()
    ws = wb.active
    ws.title = "Security_Report"

    #헤더 설정
    headers = ["순번", "제목", "링크", "위험도"]
    ws.append(headers)

    header_fill = PatternFill(start_color="333333", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF",bold=True)
        cell.fill = header_fill

    #데이터 추가
    for num,news in enumerate(news_list,1):
        ws.append([num,news["title"],news["link"],news["risk"]])

        if news["risk"] == "High":
            for cell in ws[ws.max_row]:
                cell.font = Font(color="FF0000", bold=True)

        
    #당일날짜 포함하여 파일명 설정 및 저장
    excel_path = Path(f"Security_Report_{date}.xlsx").resolve()
    wb.save(f"Security_Report_{date}.xlsx")
    print("엑셀파일 생성 성공")

    return excel_path

def send_email_report(news_list, excel_path, date):
    '''
    수집된 뉴스리스트중 위험도가 "high"인 뉴스와 엑셀파일을 이메일로 전송

    Args :  news_list: 수집된 뉴스리스트
            excel_path: 생성된 엑셀 주소
            date : 오늘날짜
    '''
    user_email = os.getenv("GOOGLE_EMAIL_USER")
    to_email = os.getenv("TO_EMAIL")
    app_password = os.getenv("GOOGLE_EMAIL_PASS")
    smtp_server = os.getenv("GOOGLE_SMTP_SERVER")
    smtp_port = int(os.getenv("GOOGLE_SMTP_PORT"))

    if not all([user_email,to_email,app_password,smtp_server]):
        raise ValueError("이메일 관련 환경변수가 누락되었거나 비어 있습니다.")

    #email 구성
    message = MIMEMultipart()
    message["Subject"] = "[보안 뉴스] 보안 뉴스 리포트"
    message["From"] = user_email
    message["To"] = to_email

    #high 기사만 추출
    high_news = [
        news for news in news_list if news["risk"] =="High"
    ]
    rows = []

    #HTML 표의 행 구성
    for num, news in enumerate(high_news, 1):
        title = escape(news["title"])
        link = escape(news["link"], quote=True)

        rows.append(f"""
            <tr>
                <td>{num}</td>
                <td><a href="{link}">{title}</a></td>
                <td style="color: red; font-weight: bold;">High</td>
            </tr>
        """)

    # HTML 본문 구성
    html = f"""
    <html>
        <body>
            <h2>위험도 High 보안 뉴스</h2>
            <p>총 {len(high_news)}건의 기사가 있습니다.</p>

            <table border="1" cellpadding="8" cellspacing="0"
                   style="border-collapse: collapse;">
                <tr style="background-color: #333333; color: white;">
                    <th>순번</th>
                    <th>뉴스 제목</th>
                    <th>위험도</th>
                </tr>
                {"".join(rows)}
            </table>
        </body>
    </html>
    """
    message.attach(MIMEText(html,"html","utf-8"))


    #엑셀파일 첨부
    with open(excel_path,"rb") as f:
        attachement = MIMEApplication(
            f.read(),
            _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    attachement.add_header(
        "Content-Disposition",
        "attachment",
        filename=excel_path.name
    )
    message.attach(attachement)


    #서버 접속 및 발송
    try:
        with smtplib.SMTP(smtp_server, smtp_port,local_hostname="localhost",timeout=10) as server:
            server.starttls()  # 메일 서버와의 통신을 암호화
            server.login(user_email, app_password)
            server.send_message(message)
            print(f"[{user_email}] 메일 발송 성공!")
    except Exception as e:
        print(f"오류 발생: {e}")


def main():
    date = datetime.now().strftime("%Y%m%d")

    try:
        news_list = get_security_news()

        if not news_list:
            return

        excel_path = create_excel_report(news_list, date)
        if not excel_path:
            return
        
        send_email_report(news_list, excel_path, date)

    except (OSError, ValueError, TypeError) as e:
        print(f"❌ 리포트 처리 중 오류가 발생했습니다: {e}")

              
if __name__ == "__main__":
    main()
